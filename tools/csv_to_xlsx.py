"""Turn a pathfinding benchmark CSV into an .xlsx with real numbers.

Usage: python csv_to_xlsx.py <file.csv> [more.csv ...]
Needs openpyxl (pip install openpyxl).
Writes <file>.xlsx next to each CSV. Numbers are stored as numbers, so a
Polish Excel shows decimal commas and can sort, sum and chart them.
"""
import csv
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

INT_COLUMNS = {"found", "expanded", "pushed"}
FLOAT_COLUMNS = {"cost", "avg_ms", "max_ms", "reset_ms"}

HEADER_FILL = PatternFill("solid", fgColor="0F6660")
GROUP_FILL = PatternFill("solid", fgColor="E3F1EF")


def convert(csv_path: Path) -> Path:
    with csv_path.open(newline="", encoding="utf-8") as f:
        rows = list(csv.reader(f))
    header, data = rows[0], rows[1:]

    wb = Workbook()
    ws = wb.active
    ws.title = csv_path.stem.replace("pathfinding_", "")[:31]

    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    previous_route, shade = None, False
    for values in data:
        converted = []
        for name, value in zip(header, values):
            if name in INT_COLUMNS:
                converted.append(int(value))
            elif name in FLOAT_COLUMNS:
                converted.append(float(value))
            else:
                converted.append(value)
        ws.append(converted)

        # shade every other route, so its four modes read as one block
        if values[0] != previous_route:
            shade = not shade
            previous_route = values[0]
        if shade:
            for cell in ws[ws.max_row]:
                cell.fill = GROUP_FILL

    for col, name in enumerate(header, start=1):
        letter = get_column_letter(col)
        if name == "cost":
            fmt = "0.000000"
        elif name in FLOAT_COLUMNS:
            fmt = "0.000000"
        elif name in INT_COLUMNS:
            fmt = "0"
        else:
            fmt = None
        if fmt:
            for cell in ws[letter][1:]:
                cell.number_format = fmt
        width = max(len(str(c.value)) for c in ws[letter]) + 3
        ws.column_dimensions[letter].width = max(width, 10)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    out = csv_path.with_suffix(".xlsx")
    wb.save(out)
    return out


if __name__ == "__main__":
    for arg in sys.argv[1:]:
        print("written:", convert(Path(arg)))
